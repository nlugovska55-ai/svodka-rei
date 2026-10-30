#!/usr/bin/env python3
"""Импорт таблицы гарантов Лены (Google-таблица, только чтение) в снимок для пульта.

    python3 -I scripts/import_garanty.py <файл.xlsx | .csv | URL выгрузки csv> --url <ссылка> --title <название> [--tab <вкладка>] [--gid N] [--live]

Живую таблицу «Утро» скачивает коннектором Google Drive (download_file_content, xlsx или csv) и передаёт файл сюда.
Без --live пульт пишет «живая таблица не подключена» и показывает строки как последний известный файл.

Пишет /mnt/project-files/daily-control/garanty-leny/snapshot.json (строки как в таблице + разбор)
и import-status.json (последний успешный импорт, последняя попытка, ошибка).
Ошибка или пустая таблица НЕ затирают прежний снимок: пульт показывает старые строки с пометкой «устарело / сбой».
В таблицу Лены ничего не пишется.
"""
import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import urllib.request
from datetime import datetime, timedelta, timezone

MSK = timezone(timedelta(hours=3))
DIR = os.environ.get("GARANTY_DIR", "/mnt/project-files/daily-control/garanty-leny")

# Колонка → слова в заголовке. Подходит шаблон Лены «Оплаты и зарплата — <месяц>» (№, Менеджер, Клиент, Услуга, Договор, Прил.,
# Прогноз, Гарант, Оплачено, Дата план., Дата оплаты, Комментарий) и файл «октябрь 2026» (наименование клиента, услуга,
# номер договора, ожидаемая сумма, сумма пришедшая). Сумма в шаблоне «переезжает вправо»: обсуждаем → гарант → деньги пришли.
COLS = [
    ("datePaid", r"дата оплат"),
    ("datePlan", r"дата план|срок|обещ|когда"),
    ("client", r"клиент|контрагент|наименован|организац|заказчик"),
    ("manager", r"менеджер|ответствен"),
    ("service", r"услуг"),
    ("contract", r"договор"),
    ("app", r"^прил"),
    ("forecast", r"прогноз"),
    ("guarantee", r"гарант"),
    ("paid", r"оплачено|пришед|пришл|получен|поступ"),
    ("expected", r"ожида|сумм"),
    ("comment", r"коммент|примеч"),
]
STAGE = {"forecast": "обсуждаем (прогноз)", "guarantee": "клиент подтвердил (гарант)", "expected": "ожидается", "paid": "деньги пришли, по таблице"}


def now():
    return datetime.now(MSK).strftime("%Y-%m-%dT%H:%M:%S+03:00")


def num(v):
    """«12 000», «12,000 ₽», «12000,50» → число; пусто или текст → None (не 0)."""
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").replace(" ", " ").replace("₽", "").strip()
    if not re.search(r"\d", s):
        return None
    s = re.sub(r"(?<=\d)[ ,](?=\d{3}(?!\d))", "", s).replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None


def ymd(v, year):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return s[:10]
    m = re.match(r"^(\d{1,2})\.(\d{1,2})(?:\.(\d{2,4}))?", s)
    if not m:
        return None
    y = m.group(3) or str(year)
    return f"{'20' + y if len(y) == 2 else y}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"


def header_row(rows):
    for i, r in enumerate(rows[:20]):
        line = " ".join(str(c) for c in r).lower()
        if re.search(r"клиент", line) and re.search(r"прогноз|гарант|оплачено|ожида|сумм", line):
            return i
    raise ValueError("не найден заголовок: нужна строка с колонкой клиента и суммами (прогноз/гарант/оплачено или ожидаемая)")


def colmap(head):
    used, out = set(), {}
    for key, rx in COLS:
        for j, h in enumerate(head):
            if j not in used and re.search(rx, str(h).lower().strip()):
                out[key] = j
                used.add(j)
                break
    if "client" not in out or not ({"forecast", "guarantee", "expected", "paid"} & set(out)):
        raise ValueError(f"в заголовке нет колонок клиента и сумм: {head}")
    return out


def parse_rows(rows, year):
    h = header_row(rows)
    head = [str(c).strip() for c in rows[h]]
    cm = colmap(head)
    out = []
    for n, r in enumerate(rows[h + 1:], start=h + 2):
        r = list(r) + [""] * (len(head) - len(r))
        cell = lambda k: r[cm[k]] if k in cm else ""
        txt = lambda k: str(cell(k) if cell(k) is not None else "").strip()
        if any(re.match(r"^\s*итого", str(c or "").lower()) for c in r):
            continue
        sums = {k: num(cell(k)) for k in ("forecast", "guarantee", "expected", "paid")}
        if not txt("client") or not any(v for v in sums.values()):
            continue
        stage = next((k for k in ("guarantee", "forecast", "expected") if sums[k]), "paid")
        amount = sums[stage]
        doc = " ".join(x for x in (txt("contract"), ("прил. " + txt("app")) if txt("app") else "") if x)
        rec = {"row": n, "client": txt("client"), "service": txt("service"), "doc": doc, "amount": amount,
               "amountRaw": txt(stage), "stage": stage, "status": STAGE[stage],
               "date": ymd(cell("datePlan"), year), "dateRaw": txt("datePlan"), "datePaid": ymd(cell("datePaid"), year),
               "received": sums["paid"], "receivedRaw": txt("paid"), "who": txt("manager"), "note": txt("comment"),
               "raw": {hh: str(v if v is not None else "").strip() for hh, v in zip(head, r) if hh}}
        rec["key"] = hashlib.sha1(f"{compact(rec['client'])}|{compact(doc)}|{compact(rec['service'])}|{amount}".encode()).hexdigest()[:10]
        out.append(rec)
    if not out:
        raise ValueError("в таблице нет строк с суммами")
    return head, out


def read_xlsx(path, tab):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = [tab] if tab else wb.sheetnames
    for nm in names:
        rows = [list(r) for r in wb[nm].iter_rows(values_only=True)]
        try:
            header_row(rows)
            return nm, rows
        except ValueError:
            continue
    raise ValueError(f"ни на одной вкладке ({', '.join(names)}) нет шапки таблицы гарантов")


def compact(s):
    return "".join(re.findall(r"[a-zа-яё0-9]+", str(s).lower().replace("ё", "е")))


def save(path, data):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, path)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("src", help="csv-файл или URL выгрузки csv")
    ap.add_argument("--url", default="", help="ссылка на таблицу для пульта (исходная строка)")
    ap.add_argument("--tab", default="", help="название вкладки")
    ap.add_argument("--gid", default="", help="gid вкладки для ссылки на строку")
    ap.add_argument("--title", default="", help="название таблицы")
    ap.add_argument("--live", action="store_true", help="это живая таблица Лены (не копия и не вложение из письма)")
    a = ap.parse_args()
    os.makedirs(DIR, exist_ok=True)
    sp, stp = os.path.join(DIR, "snapshot.json"), os.path.join(DIR, "import-status.json")
    st = json.load(open(stp, encoding="utf-8")) if os.path.exists(stp) else {}
    st["lastAttempt"] = now()
    try:
        tab = a.tab
        if re.match(r"^https?://", a.src):
            with urllib.request.urlopen(a.src, timeout=30) as r:
                text = r.read().decode("utf-8-sig")
            if text.lstrip().startswith("<"):
                raise ValueError("вместо таблицы пришла страница входа: таблица не открыта по ссылке")
            rows = list(csv.reader(io.StringIO(text)))
        elif a.src.lower().endswith(".xlsx"):
            tab, rows = read_xlsx(a.src, a.tab)
            text = json.dumps(rows, ensure_ascii=False, default=str)
        else:
            text = open(a.src, encoding="utf-8-sig").read()
            rows = list(csv.reader(io.StringIO(text)))
        head, recs = parse_rows(rows, datetime.now(MSK).year)
    except Exception as e:  # прежний снимок остаётся как есть
        st.update({"ok": False, "lastError": f"{type(e).__name__}: {e}"})
        save(stp, st)
        print("ОШИБКА импорта, прежний снимок сохранён:", st["lastError"], file=sys.stderr)
        sys.exit(1)
    snap = {"importedAt": st["lastAttempt"], "source": {"url": a.url, "tab": tab, "live": a.live, "title": a.title, "gid": a.gid, "file": a.src if not a.src.startswith("http") else ""},
            "header": head, "rows": recs}
    save(sp, snap)
    st.update({"ok": True, "lastOk": st["lastAttempt"], "lastError": "", "rows": len(recs), "live": a.live,
               "sha": hashlib.sha1(text.encode()).hexdigest()[:12]})
    save(stp, st)
    print(f"гаранты: {len(recs)} строк -> {sp}")


if __name__ == "__main__":
    main()
